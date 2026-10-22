"""Assemble the StaticRecord: one full-bleed screenshot per slide, each
slide's speaker notes written into the .pptx notes (Part 9).

Run after:  node tools/capture_slides.mjs     (export/slides/*.png + notes.json)
Then:       python3 tools/build_pptx.py        -> export/CS667_Technical_Presentation.pptx

Re-opens the file and asserts screenshots == slides == notes.
"""
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "export" / "slides"
OUT = ROOT / "export" / "CS667_Technical_Presentation.pptx"

DISCLAIMER = ("This .pptx is the static record of the presentation; the HTML deck (cs667-deck/index.html) "
              "is the live version. The animations, the threshold slider and the 3D map interaction exist "
              "only in the live deck.")


def main():
    shots = sorted(SHOTS.glob("slide-*.png"))
    notes = json.loads((SHOTS / "notes.json").read_text())
    if len(shots) != len(notes):
        sys.exit(f"{len(shots)} screenshots but {len(notes)} notes: re-run capture_slides.mjs")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for shot, note in zip(shots, notes):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(shot), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
        text = note["notes"]
        if not text:
            sys.exit(f"slide {note['index'] + 1} has no notes")
        slide.notes_slide.notes_text_frame.text = text
    # the closing slide's notes always carry the disclaimer (INV-11)
    closing = next(i for i, n in enumerate(notes) if n["id"] == "close")
    tf = prs.slides[closing].notes_slide.notes_text_frame
    if "static record" not in tf.text:
        tf.text = tf.text + "\n\n" + DISCLAIMER
    prs.core_properties.title = "Predicting NYC Restaurant Inspection Grades — technical presentation"
    prs.core_properties.author = "Rafiul Haider, Shadman Rahman"
    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)

    check = Presentation(OUT)
    n_slides = len(check.slides)
    n_notes = sum(1 for s in check.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
    assert len(shots) == n_slides == n_notes, (len(shots), n_slides, n_notes)
    print(f"wrote {OUT.relative_to(ROOT)}: {n_slides} slides, {n_notes} with notes, {OUT.stat().st_size / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
